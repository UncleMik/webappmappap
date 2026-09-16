from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "БРИФ_ДЛЯ_ВЛАДЕЛЬЦА_ШКОЛЫ.md"
OUTPUT = ROOT / "БРИФ_ДЛЯ_ВЛАДЕЛЬЦА_ШКОЛЫ.docx"

# Resolved preset: compact_reference_guide.
TOKENS = {
    "font": "Calibri",
    "body_size": 11,
    "body_after": 6,
    "body_line": 1.25,
    "h1_size": 16,
    "h1_before": 18,
    "h1_after": 10,
    "h2_size": 13,
    "h2_before": 14,
    "h2_after": 7,
    "h3_size": 12,
    "h3_before": 10,
    "h3_after": 5,
    "blue": "2E74B5",
    "dark_blue": "1F4D78",
    "ink": "0B2545",
    "muted": "667085",
    "table_fill": "E8EEF5",
    "light_fill": "F4F6F9",
    "border": "CBD5E1",
    "content_dxa": 9360,
    "table_indent_dxa": 120,
}


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_bottom_border(paragraph, color="D0D5DD", size="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = TOKENS["font"]
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), TOKENS["font"])
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), TOKENS["font"])
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), TOKENS["font"])
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline_markup(paragraph, text, base_size=None, base_color=None):
    cursor = 0
    for match in re.finditer(r"\*\*(.+?)\*\*|`(.+?)`|\*(.+?)\*", text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=base_size, color=base_color)
        value = match.group(1) or match.group(2) or match.group(3)
        run = paragraph.add_run(value)
        set_run_font(
            run,
            size=base_size,
            color=base_color,
            bold=match.group(1) is not None,
            italic=match.group(3) is not None,
        )
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=base_size, color=base_color)


def create_numbering(document):
    numbering = document.part.numbering_part.element

    def add_abstract(abstract_id, fmt, text, left, hanging, font=None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl.append(suff)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(left))
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), str(hanging))
        p_pr.append(ind)
        lvl.append(p_pr)
        if font:
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), font)
            r_fonts.set(qn("w:hAnsi"), font)
            r_pr.append(r_fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

        num_id = abstract_id + 100
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)
        return num_id

    bullet_id = add_abstract(80, "bullet", "•", 540, 280, TOKENS["font"])
    number_id = add_abstract(81, "decimal", "%1.", 540, 280, TOKENS["font"])
    return bullet_id, number_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)


def add_num_instance(document, abstract_id, num_id):
    numbering = document.part.numbering_part.element
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Страница ")
    set_run_font(run, size=9, color=TOKENS["muted"])
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = TOKENS["font"]
    normal._element.rPr.rFonts.set(qn("w:ascii"), TOKENS["font"])
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), TOKENS["font"])
    normal.font.size = Pt(TOKENS["body_size"])
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(TOKENS["body_after"])
    normal.paragraph_format.line_spacing = TOKENS["body_line"]
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in (
        ("Heading 1", TOKENS["h1_size"], TOKENS["blue"], TOKENS["h1_before"], TOKENS["h1_after"]),
        ("Heading 2", TOKENS["h2_size"], TOKENS["blue"], TOKENS["h2_before"], TOKENS["h2_after"]),
        ("Heading 3", TOKENS["h3_size"], TOKENS["dark_blue"], TOKENS["h3_before"], TOKENS["h3_after"]),
    ):
        style = styles[name]
        style.font.name = TOKENS["font"]
        style._element.rPr.rFonts.set(qn("w:ascii"), TOKENS["font"])
        style._element.rPr.rFonts.set(qn("w:hAnsi"), TOKENS["font"])
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = hp.add_run("Школа родов и материнства  |  Продуктовый бриф")
    set_run_font(run, size=9, color=TOKENS["muted"], bold=True)
    hp.paragraph_format.space_after = Pt(2)
    add_bottom_border(hp, color="D7DEE8", size="4")

    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_field(fp)


def add_opening(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("ПРОДУКТОВЫЙ БРИФ")
    set_run_font(run, size=10, color=TOKENS["blue"], bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("Веб-приложение школы родов и материнства")
    set_run_font(run, size=25, color=TOKENS["ink"], bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("Анкета для владельца школы: пользовательский продукт, админ-панель и границы MVP")
    set_run_font(run, size=12.5, color=TOKENS["muted"])

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [3120, 3120, 3120], indent=0)
    set_repeat_table_header(table.rows[0])
    for cell, label, value in zip(
        table.rows[0].cells,
        ("ПРОХОД 1", "ПРОХОД 2", "ПРОХОД 3"),
        ("Ключевые решения", "Пользовательский опыт", "Админ-панель"),
    ):
        shade_cell(cell, TOKENS["light_fill"])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label + "\n")
        set_run_font(r, size=9, color=TOKENS["blue"], bold=True)
        r = p.add_run(value)
        set_run_font(r, size=10.5, color=TOKENS["ink"], bold=True)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def add_option(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.widow_control = True
    box = p.add_run("☐  ")
    set_run_font(box, size=12, color=TOKENS["blue"], bold=True)
    add_inline_markup(p, text, base_size=10.8)


def add_response_field(doc, label, lines=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    add_inline_markup(p, label, base_size=10.5, base_color=TOKENS["dark_blue"])
    for _ in range(lines):
        line = doc.add_paragraph()
        line.paragraph_format.space_before = Pt(0)
        line.paragraph_format.space_after = Pt(5)
        line.paragraph_format.line_spacing = 1
        line.add_run(" ")
        add_bottom_border(line, color="C9D1DC", size="4")


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [TOKENS["content_dxa"]])
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    shade_cell(cell, TOKENS["light_fill"])
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    marker = p.add_run("ВАЖНО  ")
    set_run_font(marker, size=9.5, color=TOKENS["blue"], bold=True)
    add_inline_markup(p, text, base_size=10.5, base_color=TOKENS["ink"])
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_markdown_table(doc, rows):
    if not rows:
        return
    parsed = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows]
    if len(parsed) > 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in parsed[1]):
        parsed.pop(1)
    cols = len(parsed[0])
    table = doc.add_table(rows=len(parsed), cols=cols)
    table.style = "Table Grid"
    if cols == 6:
        widths = [3210, 1230, 1230, 1230, 1230, 1230]
    elif cols == 5:
        widths = [4680, 1170, 1170, 1170, 1170]
    else:
        base = TOKENS["content_dxa"] // cols
        widths = [base] * cols
        widths[-1] += TOKENS["content_dxa"] - sum(widths)
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for r_idx, values in enumerate(parsed):
        for c_idx, value in enumerate(values):
            cell = table.cell(r_idx, c_idx)
            if r_idx == 0:
                shade_cell(cell, TOKENS["table_fill"])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.08
            add_inline_markup(
                p,
                value,
                base_size=9.2 if cols >= 5 else 10,
                base_color=TOKENS["ink"],
            )
            for run in p.runs:
                if r_idx == 0:
                    run.bold = True
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(4)


def build():
    markdown = SOURCE.read_text(encoding="utf-8")
    lines = markdown.splitlines()
    doc = Document()
    configure_document(doc)
    bullet_id, number_id = create_numbering(doc)
    add_opening(doc)

    first_title_skipped = False
    blank_sequence_num_id = None
    next_blank_num_id = 300
    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_markdown_table(doc, table_lines)
            continue
        if line == "---":
            index += 1
            continue
        if line.startswith("# ") and not first_title_skipped:
            first_title_skipped = True
            index += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.page_break_before = False
            add_inline_markup(p, line[2:], base_size=TOKENS["h1_size"], base_color=TOKENS["blue"])
            index += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline_markup(p, line[3:], base_size=TOKENS["h1_size"], base_color=TOKENS["blue"])
            index += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            p.paragraph_format.keep_with_next = True
            add_inline_markup(p, line[4:], base_size=TOKENS["h2_size"], base_color=TOKENS["blue"])
            index += 1
            continue
        if line.startswith("> "):
            add_callout(doc, line[2:])
            index += 1
            continue
        option = re.match(r"^- \[ \] (.+)$", line)
        if option:
            add_option(doc, option.group(1))
            index += 1
            continue
        bullet = re.match(r"^- (.+)$", line)
        if bullet:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = TOKENS["body_line"]
            apply_num(p, bullet_id)
            add_inline_markup(p, bullet.group(1), base_size=10.8)
            index += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = TOKENS["body_line"]
            apply_num(p, number_id)
            add_inline_markup(p, numbered.group(1), base_size=10.8)
            index += 1
            continue
        blank_numbered = re.match(r"^\d+\.$", line)
        if blank_numbered:
            item_number = int(line[:-1])
            if item_number == 1 or blank_sequence_num_id is None:
                blank_sequence_num_id = next_blank_num_id
                add_num_instance(doc, 81, blank_sequence_num_id)
                next_blank_num_id += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(6)
            apply_num(p, blank_sequence_num_id)
            run = p.add_run(" ")
            set_run_font(run, size=10.8)
            add_bottom_border(p, color="C9D1DC", size="4")
            index += 1
            continue
        if line in {"**Свой вариант:**", "**Имя/роль ответственного:**", "**Что здесь нужно исправить или дополнить:**  ", "**Что здесь нужно исправить или дополнить:**"}:
            add_response_field(doc, line.rstrip(), 1)
            index += 1
            continue
        if line == "Ответ:":
            add_response_field(doc, "Ответ:", 2)
            index += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(TOKENS["body_after"])
        p.paragraph_format.line_spacing = TOKENS["body_line"]
        add_inline_markup(p, line, base_size=10.8)
        index += 1

    props = doc.core_properties
    props.title = "Продуктовый бриф для веб-приложения школы родов и материнства"
    props.subject = "Анкета владельца школы: пользовательское приложение, админ-панель и MVP"
    props.author = "Команда продукта"
    props.keywords = "продуктовый бриф, беременность, материнство, веб-приложение, MVP"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
